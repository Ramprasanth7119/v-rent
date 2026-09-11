"""Turn the operations manual into a Word document.

Straight conversion of the Markdown: headings, paragraphs, bullets, numbered
lists and tables, with a cover page and a repeating footer. Deliberately plain —
the point of the document is that somebody can read it, not that it is designed.
"""
import io
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

SRC = r'C:\Users\hp\Desktop\v-rent\docs\OPERATIONS-MANUAL.md'
OUT = sys.argv[1] if len(sys.argv) > 1 else r'C:\Users\hp\Desktop\v-rent\docs\V-RENT-Operations-Manual.docx'

INK = RGBColor(0x12, 0x23, 0x2A)
MUTED = RGBColor(0x5A, 0x6E, 0x76)
ACCENT = RGBColor(0x0A, 0x5C, 0x73)


def shade(cell, hex_colour):
    tc = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:fill'), hex_colour)
    tc.append(el)


def footer_text(section, text):
    p = section.footer.paragraphs[0]
    p.text = text
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED


def inline(paragraph, text):
    """Bold **runs** and monospace `runs`; everything else is plain."""
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = paragraph.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
        else:
            paragraph.add_run(part)


doc = Document()

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)
style.font.color.rgb = INK
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

for name, size, colour, before in (
    ('Heading 1', 20, INK, 18),
    ('Heading 2', 14, INK, 14),
    ('Heading 3', 11.5, ACCENT, 10),
):
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = colour
    st.font.bold = True
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(4)

for s in doc.sections:
    s.top_margin = Inches(0.9)
    s.bottom_margin = Inches(0.9)
    s.left_margin = Inches(0.95)
    s.right_margin = Inches(0.95)
    footer_text(s, 'V-RENT — operations manual')

# ------------------------------------------------------------------ cover
title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(150)
r = title.add_run('V-RENT')
r.font.size = Pt(40)
r.font.bold = True
r.font.color.rgb = INK

sub = doc.add_paragraph()
r = sub.add_run('Operations manual')
r.font.size = Pt(18)
r.font.color.rgb = ACCENT

blurb = doc.add_paragraph()
r = blurb.add_run(
    'A guide for somebody opening V-RENT for the first time. Every section says where to click '
    'and what you get. Nothing here needs a technical background.'
)
r.font.size = Pt(11)
r.font.color.rgb = MUTED
blurb.paragraph_format.space_before = Pt(14)

doc.add_page_break()

# ------------------------------------------------------------------- body
lines = io.open(SRC, encoding='utf-8').read().split('\n')
i = 0
skipped_title = False

while i < len(lines):
    line = lines[i].rstrip()

    if line.startswith('# '):
        # The cover already carries the document title.
        if not skipped_title:
            skipped_title = True
            i += 1
            # Skip the blurb that follows it, which the cover also carries.
            while i < len(lines) and not lines[i].startswith('#'):
                i += 1
            continue
        doc.add_heading(line[2:], level=1)
        i += 1
        continue

    if line.startswith('### '):
        doc.add_heading(line[4:], level=3)
        i += 1
        continue

    if line.startswith('## '):
        doc.add_heading(line[3:], level=2)
        i += 1
        continue

    if line.startswith('---'):
        i += 1
        continue

    # Table
    if line.startswith('|') and i + 1 < len(lines) and set(lines[i + 1].replace('|', '').strip()) <= set('- :'):
        header = [c.strip() for c in line.strip('|').split('|')]
        i += 2
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            rows.append([c.strip() for c in lines[i].strip('|').split('|')])
            i += 1
        table = doc.add_table(rows=1, cols=len(header))
        table.style = 'Table Grid'
        for n, h in enumerate(header):
            cell = table.rows[0].cells[n]
            cell.text = ''
            para = cell.paragraphs[0]
            run = para.add_run(re.sub(r'\*\*|`', '', h))
            run.bold = True
            run.font.size = Pt(9.5)
            shade(cell, 'F2F6F7')
        for row in rows:
            cells = table.add_row().cells
            for n, value in enumerate(row[:len(header)]):
                cells[n].text = ''
                para = cells[n].paragraphs[0]
                inline(para, value)
                for run in para.runs:
                    run.font.size = Pt(9.5)
        doc.add_paragraph()
        continue

    def take_continuations(start):
        """A list item wrapped over several lines is still one item."""
        buf = []
        n = start
        while n < len(lines) and lines[n].startswith('  ') and lines[n].strip():
            buf.append(lines[n].strip())
            n += 1
        return buf, n

    # Numbered list
    m = re.match(r'^(\d+)\.\s+(.*)', line)
    if m:
        rest, i = take_continuations(i + 1)
        para = doc.add_paragraph(style='List Number')
        inline(para, ' '.join([m.group(2)] + rest))
        continue

    # Bullet
    if line.startswith('- '):
        rest, i = take_continuations(i + 1)
        para = doc.add_paragraph(style='List Bullet')
        inline(para, ' '.join([line[2:]] + rest))
        continue

    if not line.strip():
        i += 1
        continue

    # Paragraph, joining any wrapped continuation lines.
    buf = [line]
    i += 1
    while i < len(lines) and lines[i].strip() and not re.match(r'^([#|\-\d]|\s*$)', lines[i]):
        buf.append(lines[i].strip())
        i += 1
    para = doc.add_paragraph()
    inline(para, ' '.join(buf))

doc.save(OUT)
print('wrote', OUT)
